"""The plan as files you can take away from the app.

One list of cards feeds the PDF, the slideshow and the audio, so the three
cannot disagree about what a step says — and the audio is named from the same
card as the slide that mentions it. These tests pin that agreement, because
it is invisible until a slide points at a clip that is not there.
"""
from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from corpus.sentence import Sentence
from deck import Card, Example, cards_from, slug
from deck.speech import lines_for, speak, split_gloss, write_manifest
from roadmap.step import RoadmapStep
from vocab.entry import Unit


def step(position: int, key: str, text: str, translation=None,
         pattern: bool = False, beside: str | None = None) -> RoadmapStep:
    return RoadmapStep(
        position=position,
        unit=Unit.pattern(key) if pattern else Unit.exact(key),
        sentence=Sentence(text=text, translation=translation),
        gain=0, score=0.0, now_readable=0,
        beside=Unit.exact(beside) if beside else None,
    )


PLAN = [
    step(1, "genau", "Ich will genau wissen, was du an ihm findest."),
    step(2, "die Zeit", "Ich gebe ihr etwas Zeit.", "I give her some time",
         pattern=True),
    step(3, "gehen", "Wir gehen jetzt.", beside="jetzt"),
]


class CardTest(unittest.TestCase):
    def test_a_step_becomes_a_card(self) -> None:
        first = cards_from(PLAN)[0]
        self.assertEqual(first.position, 1)
        self.assertEqual(first.word, "genau")
        self.assertFalse(first.is_pattern)
        self.assertIsNone(first.translation)

    def test_a_pattern_says_so(self) -> None:
        self.assertTrue(cards_from(PLAN)[1].is_pattern)

    def test_a_translation_is_carried(self) -> None:
        self.assertEqual(cards_from(PLAN)[1].translation, "I give her some time")

    def test_a_relaxed_step_carries_its_second_word(self) -> None:
        """Not i+1, and every renderer has to be able to say so."""
        self.assertEqual(cards_from(PLAN)[2].beside, "jetzt")
        self.assertIsNone(cards_from(PLAN)[0].beside)

    def test_every_card_knows_the_length_of_the_plan(self) -> None:
        self.assertEqual({c.total for c in cards_from(PLAN)}, {3})

    def test_the_stem_carries_the_order_and_the_word(self) -> None:
        """A folder of five-digit numbers says nothing about what you are
        scrubbing through."""
        card = Card(412, "die Geschichte", True, (Example("Satz."),), None, 3861)
        self.assertEqual(card.stem, "00412-die-Geschichte")

    def test_the_stem_is_padded(self) -> None:
        """A media player goes by the name, and two steps can teach words
        that sort the other way round from their positions."""
        self.assertTrue(cards_from(PLAN)[0].stem.startswith("00001-"))
        self.assertEqual(
            Card(412, "x", False, (Example("y"),), None, 3861).stem, "00412-x")
        # Padded, so a directory listing sorts the way the plan runs.
        names = sorted(Card(n, "w", False, (Example("s"),), None, 3861).stem
                       for n in (2, 10, 1000))
        self.assertEqual([n.split("-")[0] for n in names],
                         ["00002", "00010", "01000"])


class SlugTest(unittest.TestCase):
    """A word as something safe to put in a filename."""

    def test_spaces_become_hyphens(self) -> None:
        self.assertEqual(slug("sich auf etwas freuen"),
                         "sich-auf-etwas-freuen")

    def test_umlauts_are_written_out_not_flattened(self) -> None:
        """`ü` to `u` makes `fuhren` of `führen`, which is a different word.
        German writes them this way when it cannot print them."""
        self.assertEqual(slug("führen"), "fuehren")
        self.assertEqual(slug("zurücktreten"), "zuruecktreten")
        self.assertEqual(slug("größer"), "groesser")
        self.assertEqual(slug("Öl"), "Oel")

    def test_capitals_are_kept(self) -> None:
        """German nouns carry one, and it is information."""
        self.assertEqual(slug("die Geschichte"), "die-Geschichte")

    def test_a_hyphenated_word_stays_one_word(self) -> None:
        self.assertEqual(slug("die E-Mail"), "die-E-Mail")

    def test_an_accent_loses_the_accent_not_the_letter(self) -> None:
        self.assertEqual(slug("das Café"), "das-Cafe")

    def test_a_long_key_is_cut_but_not_left_ragged(self) -> None:
        name = slug("sich mit jemandem über etwas unterhalten")
        self.assertLessEqual(len(name), 48)
        self.assertFalse(name.endswith("-"))

    def test_nothing_unsafe_survives(self) -> None:
        for word in ("etw./jdn. (Akk) lassen", "gern, gerne", "a/b\\c:d*e?f"):
            name = slug(word)
            for junk in "/\\:*?\"<>| .,()":
                self.assertNotIn(junk, name, f"{word} -> {name}")


class ManifestTest(unittest.TestCase):
    def test_it_names_the_file_the_card_owns(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = write_manifest(cards_from(PLAN), Path(tmp) / "deck.tsv")
            rows = list(csv.DictReader(open(path, encoding="utf-8"),
                                       delimiter="\t"))
        self.assertEqual(rows[0]["audio"],
                         f"[sound:{cards_from(PLAN)[0].stem}.wav]")
        self.assertEqual(rows[1]["word"], "die Zeit")

    def test_a_sentence_with_a_comma_survives(self) -> None:
        """Tabs rather than commas, which is why the corpus can say what it
        likes inside a field."""
        with tempfile.TemporaryDirectory() as tmp:
            path = write_manifest(cards_from(PLAN), Path(tmp) / "deck.tsv")
            rows = list(csv.DictReader(open(path, encoding="utf-8"),
                                       delimiter="\t"))
        self.assertEqual(rows[0]["sentence"], PLAN[0].sentence.text)


class Recorder:
    """A voice that records what it was asked to say instead of speaking."""

    def __init__(self, name: str = "recorder", rate: int = 22050) -> None:
        self.name = name
        self.sample_rate = rate
        self.said: list[tuple[str, bool]] = []

    def pcm(self, text: str, slow: bool = False) -> bytes:
        self.said.append((text, slow))
        return b"\x01\x00" * 8


def spoken(cards, de, en, tmp, **kw):
    """`speak` with the gloss gate off — these tests are about the writing.

    The gate is exercised separately in `GlossGateTest`; leaving it on here
    would make every fixture card "waiting", since the fixtures carry no
    meanings.
    """
    kw.setdefault("require_gloss", False)
    return speak(cards, de, en, tmp, **kw)


def voices(rate: int = 22050, english_rate: int | None = None):
    return (Recorder("de", rate),
            Recorder("en", english_rate if english_rate else rate))


class SpeakTest(unittest.TestCase):
    def test_one_file_per_card(self) -> None:
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            written, skipped, _ = spoken(cards_from(PLAN), de, en, Path(tmp))
            names = sorted(p.name for p in Path(tmp).glob("*.wav"))
        self.assertEqual((written, skipped), (3, 0))
        self.assertEqual(names, sorted(f"{c.stem}.wav"
                                       for c in cards_from(PLAN)))

    def test_german_and_english_go_to_different_voices(self) -> None:
        """A Piper voice speaks one language. Read by the German voice,
        "What can one learn" comes out as German phonemes wearing English
        spelling."""
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN), de, en, Path(tmp))
        german = [text for text, _ in de.said]
        english = [text for text, _ in en.said]
        self.assertIn("Ich gebe ihr etwas Zeit.", german)
        self.assertIn("I give her some time", english)
        self.assertNotIn("I give her some time", german)

    def test_the_word_is_spoken_as_it_is_read_not_as_it_is_stored(self) -> None:
        de, en = voices()
        card = Card(1, "jdm. (Dat) passieren", True,
                    (Example("Das ist passiert.", "That happened."),), None, 1)
        with tempfile.TemporaryDirectory() as tmp:
            spoken([card], de, en, Path(tmp))
        self.assertIn("jemandem passieren", [text for text, _ in de.said])

    def test_slow_repeats_only_the_teaching_sentence(self) -> None:
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN)[:1], de, en, Path(tmp), slow=True)
        slowly = [text for text, is_slow in de.said if is_slow]
        self.assertEqual(slowly, [PLAN[0].sentence.text])

    def test_voices_that_disagree_on_a_rate_are_refused(self) -> None:
        """The lines are laid end to end as raw PCM, so a mismatch would play
        one language at the wrong pitch rather than failing."""
        de, en = voices(22050, english_rate=16000)
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(SystemExit):
                spoken(cards_from(PLAN), de, en, Path(tmp))

    def test_a_second_run_does_the_ones_that_are_missing(self) -> None:
        """The deck takes a while and laptops sleep. A run that starts again
        from nothing after 3,000 files is a worse failure than a slow one."""
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN), *voices(), Path(tmp))
            (Path(tmp) / f"{cards_from(PLAN)[1].stem}.wav").unlink()
            de, en = voices()
            written, skipped, _ = spoken(cards_from(PLAN), de, en, Path(tmp))
        self.assertEqual((written, skipped), (1, 2))
        self.assertIn(PLAN[1].sentence.text, [text for text, _ in de.said])

    def test_overwrite_reads_everything_again(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN), *voices(), Path(tmp))
            written, skipped, _ = spoken(cards_from(PLAN), *voices(), Path(tmp),
                                         overwrite=True)
        self.assertEqual((written, skipped), (3, 0))

    def test_an_empty_sentence_is_not_written(self) -> None:
        """A zero-length WAV plays as silence and gets mistaken for a model
        that has stopped working."""
        cards = cards_from([step(1, "x", "   ")])
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            written, skipped, _ = spoken(cards, de, en, Path(tmp))
            self.assertEqual(list(Path(tmp).glob("*.wav")), [])
        self.assertEqual((written, skipped), (0, 1))
        self.assertEqual(de.said, [])


class SpeakLogTest(unittest.TestCase):
    def test_each_written_card_is_reported(self) -> None:
        """Counts alone make a long run opaque: you can see it moving but
        not what it is moving through."""
        de, en = voices()
        seen = []
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN), de, en, Path(tmp),
                   on_card=lambda card, written: seen.append(card.spoken))
        self.assertEqual(seen, [c.spoken for c in cards_from(PLAN)])

    def test_a_skipped_card_is_not_reported(self) -> None:
        """A second run should log what it did, not what was already done."""
        with tempfile.TemporaryDirectory() as tmp:
            spoken(cards_from(PLAN), *voices(), Path(tmp))
            seen = []
            spoken(cards_from(PLAN), *voices(), Path(tmp),
                   on_card=lambda card, written: seen.append(card.spoken))
        self.assertEqual(seen, [])


class GlossGateTest(unittest.TestCase):
    """A card with no English is left for a later run, not read in German.

    The two skip rules would otherwise fight. A German-only clip is a file
    that exists, and a file that exists is skipped forever — so every card
    read before its gloss landed would stay half a card however many times
    the audio was run again. Leaving it unread is what lets the audio run
    beside the gloss pass instead of after it.
    """

    @staticmethod
    def _glossed() -> Card:
        return Card(1, "die Zeit", True,
                    (Example("Satz.", "Sentence.", "die Zeit means time."),),
                    None, 1)

    @staticmethod
    def _bare() -> Card:
        return Card(2, "die Bank", True, (Example("Satz.", "Sentence."),),
                    None, 2)

    def test_an_ungloss_d_card_is_left_alone(self) -> None:
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            written, _, waiting = speak([self._bare()], de, en, Path(tmp))
            self.assertEqual(list(Path(tmp).glob("*.wav")), [])
        self.assertEqual((written, waiting), (0, 1))
        self.assertEqual(de.said, [])

    def test_a_glossed_card_is_read(self) -> None:
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            written, _, waiting = speak([self._glossed()], de, en, Path(tmp))
        self.assertEqual((written, waiting), (1, 0))

    def test_german_only_reads_it_anyway_when_asked(self) -> None:
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            written, _, waiting = speak([self._bare()], de, en, Path(tmp),
                                        require_gloss=False)
        self.assertEqual((written, waiting), (1, 0))

    def test_a_card_already_read_is_not_re_gated(self) -> None:
        """Once a clip exists it counts as done, gloss or no gloss — the
        gate decides what to start, not what to re-examine."""
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            speak([self._bare()], de, en, Path(tmp), require_gloss=False)
            written, skipped, waiting = speak([self._bare()], de, en,
                                              Path(tmp))
        self.assertEqual((written, skipped, waiting), (0, 1, 0))


class GlossVoiceTest(unittest.TestCase):
    """The gloss is the one bilingual line, and needs both voices.

    `etwas machen means to do something.` read whole by the English voice
    mispronounces its German half — English spelling rules applied to German
    words. The word being glossed is known, so the seam is found rather than
    guessed.
    """

    def test_the_german_head_and_english_tail_are_split(self) -> None:
        self.assertEqual(
            split_gloss("etwas machen means to do something.", "etwas machen"),
            [("etwas machen", True), ("means to do something.", False)])

    def test_the_match_ignores_case(self) -> None:
        """The model sometimes capitalises the word at the start of its
        sentence, and the text it wrote is kept rather than the one asked
        for."""
        self.assertEqual(split_gloss("Die Zeit means time.", "die Zeit"),
                         [("Die Zeit", True), ("means time.", False)])

    def test_a_gloss_that_does_not_quote_the_word_still_splits(self) -> None:
        self.assertEqual(split_gloss("laufen means to run.", "etwas laufen"),
                         [("laufen", True), ("means to run.", False)])

    def test_an_unexpected_shape_is_read_as_english(self) -> None:
        """Better spoken in one voice than dropped."""
        self.assertEqual(split_gloss("to happen to someone", "jemandem passieren"),
                         [("to happen to someone", False)])

    def test_the_word_reaches_the_german_voice_in_a_real_card(self) -> None:
        card = Card(1, "etwas machen", True,
                    (Example("Was machst du?", "What are you doing?",
                             "etwas machen means to do something."),), None, 1)
        de, en = voices()
        with tempfile.TemporaryDirectory() as tmp:
            speak([card], de, en, Path(tmp))
        self.assertIn("etwas machen", [text for text, _ in de.said])
        self.assertIn("means to do something.", [text for text, _ in en.said])
        self.assertNotIn("etwas machen means to do something.",
                         [text for text, _ in en.said])

    def test_a_gloss_is_still_said_once_per_sense(self) -> None:
        """Splitting it must not undo the collapsing."""
        card = Card(1, "die Zeit", True, tuple(
            Example(f"Satz {i}.", f"Sentence {i}.", "die Zeit means time.")
            for i in (1, 2, 3)), None, 1)
        german = [t for t, de, _, _, _ in lines_for(card) if de]
        self.assertEqual(german.count("die Zeit"), 2)   # the word, then one gloss


class RenderTest(unittest.TestCase):
    def test_the_pdf_is_a_pdf(self) -> None:
        from deck.pdf import write_pdf
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pdf(cards_from(PLAN), Path(tmp) / "d.pdf", "probe")
            head = path.read_bytes()[:5]
        self.assertEqual(head, b"%PDF-")

    def test_one_slide_per_step_after_the_title(self) -> None:
        from pptx import Presentation

        from deck.slides import write_pptx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx(cards_from(PLAN), Path(tmp) / "d.pptx", "probe")
            prs = Presentation(str(path))
        self.assertEqual(len(prs.slides), len(PLAN) + 1)

    def test_the_slide_points_at_the_clip_the_card_owns(self) -> None:
        """The one place the two halves have to agree."""
        from pptx import Presentation

        from deck.slides import write_pptx
        cards = cards_from(PLAN)
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx(cards, Path(tmp) / "d.pptx", "probe")
            prs = Presentation(str(path))
            notes = prs.slides[1].notes_slide.notes_text_frame.text
        self.assertIn(f"{cards[0].stem}.wav", notes)

    def test_the_translation_is_on_the_slide(self) -> None:
        """A card reads word, sentence, English, meaning — all four in view.

        It used to hide the English in the speaker notes, on the theory that
        a translation in sight is one you read instead of the German. The
        deck is read straight through rather than presented, so all four
        lines belong on the slide and the notes carry the audio name.
        """
        from pptx import Presentation

        from deck.slides import write_pptx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx(cards_from(PLAN), Path(tmp) / "d.pptx", "probe")
            slide = Presentation(str(path)).slides[2]
            on_slide = " ".join(s.text_frame.text for s in slide.shapes
                                if s.has_text_frame)
        self.assertIn("I give her some time", on_slide)


class SenseTest(unittest.TestCase):
    """A meaning is shown once per sense, not once per sentence.

    The model groups the examples by what the word means in them and gives
    every sentence in a group the same words, so equal strings are the signal
    that the sense has not changed. Repeating it under each sentence said the
    same thing three times; dropping it entirely would hide the words that
    genuinely shift meaning between examples.
    """

    SHIFTS = [
        step(1, "die Bank", "Ich sitze auf der Bank.", pattern=True),
    ]

    @staticmethod
    def _card(*means: str) -> Card:
        return Card(1, "die Bank", True,
                    tuple(Example(f"Satz {i}.", f"Sentence {i}.", m)
                          for i, m in enumerate(means, 1)),
                    None, 1)

    def test_one_sense_is_stated_once_in_the_pdf(self) -> None:
        from deck.pdf import write_pdf
        card = self._card("die Bank means bench.", "die Bank means bench.",
                          "die Bank means bench.")
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pdf([card], Path(tmp) / "d.pdf", "probe")
            self.assertTrue(path.exists())

    def test_a_shift_gets_its_own_line_on_the_slide(self) -> None:
        from pptx import Presentation

        from deck.slides import write_pptx
        card = self._card("die Bank means bench.", "die Bank means bench.",
                          "die Bank means the financial institution.")
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx([card], Path(tmp) / "d.pptx", "probe")
            slide = Presentation(str(path)).slides[1]
            text = " ".join(s.text_frame.text for s in slide.shapes
                            if s.has_text_frame)
        self.assertEqual(text.count("means bench"), 1)
        self.assertEqual(text.count("financial institution"), 1)

    def test_the_same_sense_is_not_repeated_on_the_slide(self) -> None:
        from pptx import Presentation

        from deck.slides import write_pptx
        card = self._card(*["die Bank means bench."] * 3)
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx([card], Path(tmp) / "d.pptx", "probe")
            slide = Presentation(str(path)).slides[1]
            text = " ".join(s.text_frame.text for s in slide.shapes
                            if s.has_text_frame)
        self.assertEqual(text.count("means bench"), 1)


if __name__ == "__main__":
    unittest.main()
